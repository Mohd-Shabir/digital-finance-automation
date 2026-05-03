"""
generate_dataset.py
--------------------
Generates a synthetic corporate finance dataset for the
Digital Finance Automation project.

Output: finance_dataset.xlsx with 3 sheets:
  - Transactions     (15,000 rows)
  - Monthly Summary  (aggregated by year/month/department)
  - Department Summary (high-level KPIs per department)

Requirements:
    pip install pandas openpyxl
"""

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import random
from datetime import datetime, timedelta

# ── Reproducibility ───────────────────────────────────────────────────────────
np.random.seed(42)
random.seed(42)

# ── Reference Data ────────────────────────────────────────────────────────────
DEPARTMENTS = [
    "Engineering", "Sales", "Marketing", "Operations",
    "Finance", "HR", "IT", "Legal"
]

COST_CENTERS = {
    "Engineering": ["ENG-001", "ENG-002", "ENG-003"],
    "Sales":       ["SAL-001", "SAL-002", "SAL-003"],
    "Marketing":   ["MKT-001", "MKT-002"],
    "Operations":  ["OPS-001", "OPS-002", "OPS-003"],
    "Finance":     ["FIN-001", "FIN-002"],
    "HR":          ["HR-001",  "HR-002"],
    "IT":          ["IT-001",  "IT-002",  "IT-003"],
    "Legal":       ["LEG-001", "LEG-002"],
}

EXPENSE_CATEGORIES = [
    "Salaries", "Travel", "Software", "Hardware", "Consulting",
    "Marketing Spend", "Office Supplies", "Training", "Utilities", "Maintenance"
]

VENDORS = [
    "Accenture", "Deloitte", "Microsoft", "AWS", "Salesforce",
    "Oracle", "SAP", "IBM", "Cisco", "Dell",
    "Adobe", "Zoom", "Slack", "Workday", "ServiceNow"
]

APPROVERS = [
    "John Smith", "Sarah Johnson", "Michael Chen", "Emma Wilson",
    "David Brown", "Lisa Garcia", "Robert Taylor", "Anna Martinez"
]

STATUS_LIST      = ["Approved", "Approved", "Approved", "Pending", "Rejected", "Under Review"]
PAYMENT_METHODS  = ["Bank Transfer", "Credit Card", "Purchase Order", "Direct Debit"]
CURRENCIES       = ["USD", "USD", "USD", "EUR", "GBP"]
NOTES_OPTIONS    = [
    "Routine expense", "Urgent procurement", "Annual contract",
    "One-time cost", "Recurring monthly", ""
]

# Annual budget per department (USD)
BUDGET_BY_DEPT = {
    "Engineering": 5_000_000,
    "Sales":       3_500_000,
    "Marketing":   2_000_000,
    "Operations":  4_000_000,
    "Finance":     1_500_000,
    "HR":          1_200_000,
    "IT":          2_500_000,
    "Legal":         800_000,
}

# ── Styling Helpers ───────────────────────────────────────────────────────────
HEADER_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HEADER_FILL  = PatternFill("solid", start_color="1F4E79")
ALT_FILL     = PatternFill("solid", start_color="EBF3FB")
CENTER       = Alignment(horizontal="center", vertical="center")
THIN_BORDER  = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def apply_header(ws, headers):
    """Write and format a header row."""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header.replace("_", " "))
        cell.font  = HEADER_FONT
        cell.fill  = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def apply_row(ws, row_idx, values, currency_cols=None, pct_cols=None):
    """Write a data row with alternating fill and optional number formats."""
    fill = ALT_FILL if row_idx % 2 == 0 else None
    currency_cols = currency_cols or []
    pct_cols      = pct_cols      or []

    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font   = Font(name="Arial", size=9)
        cell.border = THIN_BORDER
        if fill:
            cell.fill = fill
        if col_idx in currency_cols:
            cell.number_format = "#,##0.00"
        if col_idx in pct_cols:
            cell.number_format = '0.00"%"'


# ── Data Generation ───────────────────────────────────────────────────────────
def generate_transactions(n: int = 15_000) -> pd.DataFrame:
    start_date = datetime(2023, 1, 1)
    end_date   = datetime(2024, 12, 31)
    date_range = (end_date - start_date).days

    rows = []
    for i in range(n):
        dept     = random.choice(DEPARTMENTS)
        cc       = random.choice(COST_CENTERS[dept])
        cat      = random.choice(EXPENSE_CATEGORIES)
        vendor   = random.choice(VENDORS) if cat in ("Software", "Hardware", "Consulting") else "Internal"

        monthly_budget    = BUDGET_BY_DEPT[dept] / 12
        budget_allocated  = round(monthly_budget * random.uniform(0.05, 0.15), 2)
        actual_spend      = round(budget_allocated * random.gauss(1.0, 0.15), 2)
        actual_spend      = max(100.0, actual_spend)

        variance     = round(actual_spend - budget_allocated, 2)
        variance_pct = round((variance / budget_allocated) * 100, 2) if budget_allocated else 0

        date = start_date + timedelta(days=random.randint(0, date_range))

        rows.append({
            "Transaction_ID":  f"TXN-{str(i + 1).zfill(6)}",
            "Date":            date.strftime("%Y-%m-%d"),
            "Year":            date.year,
            "Month":           date.strftime("%B"),
            "Quarter":         f"Q{(date.month - 1) // 3 + 1}",
            "Department":      dept,
            "Cost_Center":     cc,
            "Expense_Category": cat,
            "Vendor":          vendor,
            "Budget_Allocated": budget_allocated,
            "Actual_Spend":    actual_spend,
            "Variance":        variance,
            "Variance_Pct":    variance_pct,
            "Status":          random.choice(STATUS_LIST),
            "Approver":        random.choice(APPROVERS),
            "Payment_Method":  random.choice(PAYMENT_METHODS),
            "Currency":        random.choice(CURRENCIES),
            "Invoice_Number":  f"INV-{random.randint(10_000, 99_999)}",
            "PO_Number":       f"PO-{random.randint(1_000, 9_999)}" if random.random() > 0.3 else "N/A",
            "Notes":           random.choice(NOTES_OPTIONS),
        })

    df = pd.DataFrame(rows).sort_values("Date").reset_index(drop=True)
    return df


def build_monthly_summary(df: pd.DataFrame) -> pd.DataFrame:
    monthly = (
        df.groupby(["Year", "Month", "Department"])
        .agg(
            Budget_Allocated  = ("Budget_Allocated", "sum"),
            Actual_Spend      = ("Actual_Spend",     "sum"),
            Transaction_Count = ("Transaction_ID",   "count"),
        )
        .reset_index()
    )
    monthly["Variance"]     = monthly["Actual_Spend"] - monthly["Budget_Allocated"]
    monthly["Variance_Pct"] = ((monthly["Variance"] / monthly["Budget_Allocated"]) * 100).round(2)
    return monthly


def build_department_summary(df: pd.DataFrame) -> pd.DataFrame:
    dept = (
        df.groupby("Department")
        .agg(
            Total_Budget      = ("Budget_Allocated", "sum"),
            Total_Actual      = ("Actual_Spend",     "sum"),
            Transaction_Count = ("Transaction_ID",   "count"),
            Avg_Transaction   = ("Actual_Spend",     "mean"),
        )
        .reset_index()
    )
    dept["Total_Variance"] = dept["Total_Actual"] - dept["Total_Budget"]
    dept["Variance_Pct"]   = ((dept["Total_Variance"] / dept["Total_Budget"]) * 100).round(2)
    dept["Avg_Transaction"] = dept["Avg_Transaction"].round(2)
    return dept


# ── Excel Writing ─────────────────────────────────────────────────────────────
def write_transactions_sheet(wb: Workbook, df: pd.DataFrame):
    ws = wb.active
    ws.title = "Transactions"

    headers = list(df.columns)
    apply_header(ws, headers)

    currency_cols = [headers.index(c) + 1 for c in ("Budget_Allocated", "Actual_Spend", "Variance") if c in headers]
    pct_cols      = [headers.index("Variance_Pct") + 1] if "Variance_Pct" in headers else []

    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        apply_row(ws, row_idx, list(row), currency_cols=currency_cols, pct_cols=pct_cols)

    col_widths = [14, 12, 6, 10, 8, 12, 12, 18, 14, 16, 14, 12, 14, 10, 14, 16, 10, 14, 12, 20]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def write_monthly_summary_sheet(wb: Workbook, df: pd.DataFrame):
    ws = wb.create_sheet("Monthly Summary")
    headers = ["Year", "Month", "Department", "Budget_Allocated", "Actual_Spend",
               "Transaction_Count", "Variance", "Variance_Pct"]
    apply_header(ws, headers)

    currency_cols = [4, 5, 7]
    pct_cols      = [8]

    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        apply_row(ws, row_idx, list(row), currency_cols=currency_cols, pct_cols=pct_cols)

    for i, w in enumerate([8, 12, 14, 18, 14, 18, 14, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"


def write_department_summary_sheet(wb: Workbook, df: pd.DataFrame):
    ws = wb.create_sheet("Department Summary")
    headers = ["Department", "Total_Budget", "Total_Actual", "Transaction_Count",
               "Avg_Transaction", "Total_Variance", "Variance_Pct"]
    apply_header(ws, headers)

    currency_cols = [2, 3, 5, 6]
    pct_cols      = [7]

    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        apply_row(ws, row_idx, list(row), currency_cols=currency_cols, pct_cols=pct_cols)

    for i, w in enumerate([16, 16, 14, 18, 16, 16, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Main ──────────────────────────────────────────────────────────────────────
def main(output_path: str = "finance_dataset.xlsx", n_rows: int = 15_000):
    print(f"Generating {n_rows:,} transactions...")
    df_transactions = generate_transactions(n_rows)

    print("Building summary sheets...")
    df_monthly = build_monthly_summary(df_transactions)
    df_dept    = build_department_summary(df_transactions)

    print("Writing Excel file...")
    wb = Workbook()
    write_transactions_sheet(wb, df_transactions)
    write_monthly_summary_sheet(wb, df_monthly)
    write_department_summary_sheet(wb, df_dept)

    wb.save(output_path)

    print(f"\n✅ Done — saved to: {output_path}")
    print(f"   Transactions:       {len(df_transactions):,} rows")
    print(f"   Monthly Summary:    {len(df_monthly):,} rows")
    print(f"   Department Summary: {len(df_dept):,} rows")


if __name__ == "__main__":
    main()
